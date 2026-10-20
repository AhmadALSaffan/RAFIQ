"""PowerPoint in Arabic, spreadsheets read and written, and pictures the model can look at."""

import io
import zipfile

import pytest
from openpyxl import load_workbook
from PIL import Image

from rafiq_agent.core.agent_runtime import build_registry, permission_key
from rafiq_agent.schemas.settings import AppSettings
from rafiq_agent.tools.image_view import ImageViewTool
from rafiq_agent.tools.sheets import ExcelReadTool, ExcelWriteTool, as_number, build_xlsx, column_stats
from rafiq_agent.tools.slides import SlidesWriteTool, build_pptx, parse_outline

OUTLINE = """# خطة الإطلاق
عرض للفريق

## الفكرة
- تطبيق للطلاب
  - الدفع عبر Apple Pay
- الإطلاق بالرياض

## الأرقام
| الشهر | الرحلات |
|---|---|
| أكتوبر | 3,400 |
"""


def test_the_outline_becomes_a_title_slide_and_one_slide_per_heading():
    title, subtitle, slides = parse_outline(OUTLINE)
    assert (title, subtitle) == ("خطة الإطلاق", "عرض للفريق")
    assert [s.title for s in slides] == ["الفكرة", "الأرقام"]
    assert slides[0].bullets == [(0, "تطبيق للطلاب"), (1, "الدفع عبر Apple Pay"), (0, "الإطلاق بالرياض")]
    assert slides[1].table == [["الشهر", "الرحلات"], ["أكتوبر", "3,400"]]


def test_the_deck_reads_right_to_left():
    with zipfile.ZipFile(io.BytesIO(build_pptx(OUTLINE, look="dark"))) as zf:
        names = set(zf.namelist())
        assert {"ppt/presentation.xml", "ppt/slides/slide1.xml", "ppt/slides/slide3.xml", "ppt/theme/theme1.xml"} <= names
        slide = zf.read("ppt/slides/slide2.xml").decode()
        table = zf.read("ppt/slides/slide3.xml").decode()
        presentation = zf.read("ppt/presentation.xml").decode()
    assert 'rtl="1"' in slide and 'algn="r"' in slide
    assert 'cs typeface="Sakkal Majalla"' in slide
    assert '<a:tblPr firstRow="1" rtl="1"/>' in table
    assert presentation.count("<p:sldId ") == 3
    # "Apple Pay" keeps its own left-to-right run inside the Arabic bullet.
    assert '<a:t>Apple Pay</a:t>' in slide


async def test_the_slides_tool_writes_inside_the_folder(tmp_path):
    tool = SlidesWriteTool(tmp_path)
    assert (await tool.run({"path": "deck", "outline": OUTLINE})).ok
    assert (tmp_path / "deck.pptx").is_file()
    assert not (await tool.run({"path": "../out.pptx", "outline": OUTLINE})).ok
    assert "3 slides" in (await tool.preview({"path": "deck.pptx", "outline": OUTLINE}) or "")


def test_numbers_and_column_stats():
    assert as_number("1,200") == 1200 and as_number("12.5%") == 0.125 and as_number("abc") is None and as_number(True) is None
    stats = column_stats(["الشهر", "المبيعات"], [["يناير", "1,000"], ["فبراير", 3000], ["مارس", ""]])
    assert stats == [{"column": "المبيعات", "count": 2, "sum": 4000, "avg": 2000, "min": 1000, "max": 3000}]


def test_a_workbook_with_numbers_formulas_rtl_and_a_chart():
    data = build_xlsx(
        [{"name": "المبيعات", "rows": [["الشهر", "المبيعات"], ["يناير", "1,200"], ["فبراير", "2,800"], ["المجموع", "=SUM(B2:B3)"]]}],
        {"type": "bar", "title": "المبيعات", "categories_column": 1, "value_columns": [2]},
    )
    ws = load_workbook(io.BytesIO(data))["المبيعات"]
    assert ws.sheet_view.rightToLeft is True
    assert ws["B2"].value == 1200 and ws["B4"].value == "=SUM(B2:B3)"
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert len(ws._charts) == 1


async def test_reading_back_what_was_written(tmp_path):
    rows = [["البند", "السعر"], ["قلم", "5"], ["دفتر", "15"]]
    assert (await ExcelWriteTool(tmp_path).run({"path": "list.xlsx", "sheets": [{"name": "Sheet", "rows": rows}]})).ok
    out = await ExcelReadTool(tmp_path).run({"path": "list.xlsx"})
    assert out.ok and "| قلم | 5 |" in out.output and "sum 20" in out.output
    (tmp_path / "a.csv").write_text("name,score\nA,1\nB,3\n", encoding="utf-8")
    assert "average 2" in (await ExcelReadTool(tmp_path).run({"path": "a.csv"})).output
    assert not (await ExcelReadTool(tmp_path).run({"path": "missing.xlsx"})).ok


async def test_the_model_can_look_at_a_picture(tmp_path):
    Image.new("RGB", (3000, 1500), "orange").save(tmp_path / "shot.png")
    out = await ImageViewTool(tmp_path).run({"path": "shot.png"})
    assert out.ok and out.output.endswith("3000×1500")
    assert out.images and out.images[0].startswith("data:image/png;base64,")
    shrunk = Image.open(io.BytesIO(__import__("base64").b64decode(out.images[0].split(",", 1)[1])))
    assert max(shrunk.size) == 1600
    assert not (await ImageViewTool(tmp_path).run({"path": "notes.txt"})).ok


@pytest.mark.parametrize("name", ["slides_write", "excel_write", "word_write"])
def test_writing_office_files_asks_like_writing_a_file(name):
    assert permission_key(name) == "filesystem_write"


async def test_office_tools_are_there_even_without_a_folder(tmp_path, monkeypatch):
    monkeypatch.setenv("RAFIQ_WORKSPACE_DIR", str(tmp_path))
    registry = await build_registry(None, AppSettings(), groups=frozenset({"files"}))
    assert {"word_write", "slides_write", "excel_read", "excel_write", "image_view"} <= registry.names()
    assert "shell_run" not in registry.names()
