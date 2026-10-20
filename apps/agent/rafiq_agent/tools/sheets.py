"""Spreadsheets: read a sheet (with a quick analysis of its numbers), and write one — Arabic
sheets right to left, a bold header, sensible column widths, formulas, and a chart.

`excel_read` takes .xlsx/.xlsm and .csv; `excel_write` makes .xlsx. Both stay inside the
task's folder like every file tool.
"""

import csv
import io
import re
from pathlib import Path
from typing import Any

from rafiq_agent.tools.base import Tool, ToolResult
from rafiq_agent.tools.filesystem import PathEscapeError, _resolve
from rafiq_agent.tools.office import mostly_arabic

MAX_ROWS_SHOWN = 60
MAX_CELLS_WRITTEN = 200_000
NUMBER = re.compile(r"^[+-]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?%?$")


def as_number(value: Any) -> float | None:
    """A cell's value as a number when it is one ("1,200", "12.5%", 7) — else None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    if isinstance(value, str) and NUMBER.match(value.strip()):
        text = value.strip().replace(",", "")
        return float(text[:-1]) / 100 if text.endswith("%") else float(text)
    return None


def column_stats(header: list[str], rows: list[list[Any]]) -> list[dict[str, Any]]:
    """For each column that's mostly numbers: count, sum, average, min and max."""
    out = []
    for c, name in enumerate(header):
        nums = [n for r in rows if c < len(r) and (n := as_number(r[c])) is not None]
        filled = sum(1 for r in rows if c < len(r) and r[c] not in (None, ""))
        if nums and len(nums) >= max(1, filled * 0.6):
            out.append({"column": name or f"#{c + 1}", "count": len(nums), "sum": sum(nums), "avg": sum(nums) / len(nums), "min": min(nums), "max": max(nums)})
    return out


def _fmt(n: float) -> str:
    return f"{n:,.0f}" if float(n).is_integer() else f"{n:,.2f}"


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("|", "\\|").replace("\n", " ")


def describe(name: str, sheets: list[tuple[str, list[list[Any]]]], sheet: str | None) -> str:
    lines = [f"{name}: {len(sheets)} sheet(s): " + ", ".join(f"{s} ({len(r)} rows)" for s, r in sheets)]
    picked = [x for x in sheets if sheet is None or x[0] == sheet] or sheets[:1]
    for title, rows in picked[:3]:
        rows = [r for r in rows if any(v not in (None, "") for v in r)]
        if not rows:
            lines.append(f"\n## {title}\n(empty)")
            continue
        width = max(len(r) for r in rows)
        rows = [list(r) + [None] * (width - len(r)) for r in rows]
        header = [_cell(v) for v in rows[0]]
        body = rows[1:]
        lines.append(f"\n## {title}  ({len(body)} data rows × {width} columns)")
        lines.append("| " + " | ".join(header) + " |")
        lines.append("|" + "---|" * width)
        for r in body[:MAX_ROWS_SHOWN]:
            lines.append("| " + " | ".join(_cell(v) for v in r) + " |")
        if len(body) > MAX_ROWS_SHOWN:
            lines.append(f"… {len(body) - MAX_ROWS_SHOWN} more rows")
        stats = column_stats(header, body)
        if stats:
            lines.append("\nNumbers:")
            for s in stats:
                lines.append(f"- {s['column']}: count {s['count']}, sum {_fmt(s['sum'])}, average {_fmt(s['avg'])}, min {_fmt(s['min'])}, max {_fmt(s['max'])}")
    return "\n".join(lines)


def read_book(target: Path) -> list[tuple[str, list[list[Any]]]]:
    if target.suffix.lower() == ".csv":
        raw = target.read_bytes()
        text = raw.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
        return [(target.stem, [list(r) for r in csv.reader(io.StringIO(text), dialect)])]
    from openpyxl import load_workbook

    book = load_workbook(target, read_only=True, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in book.worksheets]
    finally:
        book.close()


class ExcelReadTool(Tool):
    name = "excel_read"
    category = "read_only"
    description = (
        "Read a spreadsheet (.xlsx, .xlsm or .csv) relative to the task's working directory: its "
        "sheets, the rows as a table (first 60), and for every numeric column the count, sum, "
        "average, min and max. Use it before answering questions about a spreadsheet."
    )
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "sheet": {"type": "string", "description": "One sheet by name (default: the first three)"},
        },
        "required": ["path"],
    }

    def __init__(self, working_dir: Path) -> None:
        self.working_dir = working_dir

    async def run(self, args: dict[str, Any]) -> ToolResult:
        try:
            target = _resolve(self.working_dir, str(args.get("path", "")))
        except PathEscapeError as e:
            return ToolResult(ok=False, output=str(e))
        if not target.is_file():
            return ToolResult(ok=False, output=f"no such file: {args.get('path')}")
        if target.suffix.lower() not in (".xlsx", ".xlsm", ".csv"):
            return ToolResult(ok=False, output="excel_read reads .xlsx, .xlsm and .csv files")
        try:
            book = read_book(target)
        except Exception as e:  # noqa: BLE001 - a broken or locked file is reported, not raised
            return ToolResult(ok=False, output=f"couldn't read {target.name}: {e}")
        return ToolResult(ok=True, output=describe(target.name, book, args.get("sheet")))


CHARTS = ("bar", "line", "pie")


def build_xlsx(sheets: list[dict[str, Any]], chart: dict[str, Any] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="EDEFF2")
    for spec in sheets:
        ws = book.create_sheet(str(spec.get("name") or f"Sheet{len(book.worksheets) + 1}")[:31])
        rows: list[list[Any]] = spec.get("rows") or []
        text = " ".join(str(v) for r in rows for v in r if isinstance(v, str))
        rtl = spec.get("rtl") if isinstance(spec.get("rtl"), bool) else mostly_arabic(text)
        ws.sheet_view.rightToLeft = bool(rtl)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                if isinstance(value, str) and value.startswith("="):
                    cell = ws.cell(row=r, column=c, value=value)
                elif (n := as_number(value)) is not None and r > 1:
                    cell = ws.cell(row=r, column=c, value=int(n) if n.is_integer() and "%" not in str(value) else n)
                    if isinstance(value, str) and value.strip().endswith("%"):
                        cell.number_format = "0.0%"
                    elif isinstance(cell.value, int):
                        cell.number_format = "#,##0"
                    elif isinstance(cell.value, float):
                        cell.number_format = "#,##0.00"
                else:
                    cell = ws.cell(row=r, column=c, value=value)
                if r == 1:
                    cell.font = head_font
                    cell.fill = head_fill
                    cell.alignment = Alignment(horizontal="center")
        if rows:
            # A formula in a column of numbers (a total) reads like the numbers above it.
            for col in ws.iter_cols(min_row=2):
                formats = [c.number_format for c in col if isinstance(c.value, int | float)]
                if formats:
                    for c in col:
                        if isinstance(c.value, str) and c.value.startswith("="):
                            c.number_format = max(set(formats), key=formats.count)
            ws.freeze_panes = "A2"
            for c in range(1, max(len(r) for r in rows) + 1):
                longest = max((len(str(r[c - 1])) for r in rows if c - 1 < len(r) and r[c - 1] is not None), default=6)
                ws.column_dimensions[get_column_letter(c)].width = min(48, max(8, longest + 3))

    if chart and book.worksheets:
        ws = book[chart["sheet"]] if chart.get("sheet") in book.sheetnames else book.worksheets[0]
        kind = chart.get("type") if chart.get("type") in CHARTS else "bar"
        cat_col = int(chart.get("categories_column") or 1)
        value_cols = [int(c) for c in (chart.get("value_columns") or [2])][:6]
        # A totals row at the bottom (=SUM…) would dwarf every bar: the chart stops above it.
        last_row = ws.max_row
        while last_row > 2 and any(str(ws.cell(row=last_row, column=c).value or "").upper().startswith(("=SUM", "=TOTAL", "=SUBTOTAL")) for c in value_cols):
            last_row -= 1
        made = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
        made.title = chart.get("title") or None
        for col in value_cols:
            made.add_data(Reference(ws, min_col=col, min_row=1, max_row=last_row), titles_from_data=True)
        made.set_categories(Reference(ws, min_col=cat_col, min_row=2, max_row=last_row))
        made.width, made.height = 18, 9
        if kind != "pie":
            # Shown axes: newer Excel hides the ones openpyxl leaves unmarked.
            made.x_axis.delete = False
            made.y_axis.delete = False
            if ws.sheet_view.rightToLeft:
                # An Arabic sheet's chart reads from the right too.
                made.x_axis.scaling.orientation = "maxMin"
        anchor_col = get_column_letter(ws.max_column + 2)
        ws.add_chart(made, chart.get("anchor") or f"{anchor_col}2")

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class ExcelWriteTool(Tool):
    name = "excel_write"
    category = "write"
    description = (
        "Write an Excel workbook (.xlsx) relative to the task's working directory. Each sheet is "
        "a name and rows (the first row is the header, made bold and frozen). Numbers given as "
        "text (\"1,200\", \"12%\") become real numbers; text starting with = is a formula "
        "(=SUM(B2:B10)). Arabic sheets read right to left. Optionally add one chart (bar, line or "
        "pie) from a category column and value columns (1-based). Don't write .xlsx by hand."
    )
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Where to save it, ending in .xlsx"},
            "sheets": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "name": {"type": "string"},
                        "rows": {"type": "array", "items": {"type": "array", "items": {}}},
                        "rtl": {"type": "boolean", "description": "Right to left (default: from the text)"},
                    },
                    "required": ["rows"],
                },
            },
            "chart": {
                "type": "object",
                "properties": {
                    "sheet": {"type": "string"},
                    "type": {"type": "string", "enum": list(CHARTS)},
                    "title": {"type": "string"},
                    "categories_column": {"type": "integer", "description": "1-based column of the labels"},
                    "value_columns": {"type": "array", "items": {"type": "integer"}, "description": "1-based columns of the values"},
                },
            },
        },
        "required": ["path", "sheets"],
    }

    def __init__(self, working_dir: Path) -> None:
        self.working_dir = working_dir

    def _target(self, args: dict[str, Any]) -> Path:
        path = str(args.get("path", "")).strip()
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        return _resolve(self.working_dir, path)

    async def preview(self, args: dict[str, Any]) -> str | None:
        try:
            target = self._target(args)
        except PathEscapeError:
            return None
        sheets = args.get("sheets") or []
        lines = [f"{'overwrite' if target.exists() else 'new'}: {target.name}"]
        lines += [f"  {s.get('name') or 'Sheet'}: {len(s.get('rows') or [])} rows" for s in sheets if isinstance(s, dict)]
        if isinstance(args.get("chart"), dict):
            lines.append(f"  + {args['chart'].get('type', 'bar')} chart")
        return "\n".join(lines)

    async def run(self, args: dict[str, Any]) -> ToolResult:
        sheets = [s for s in (args.get("sheets") or []) if isinstance(s, dict)]
        if not sheets:
            return ToolResult(ok=False, output="no sheets given")
        cells = sum(len(r) for s in sheets for r in (s.get("rows") or []) if isinstance(r, list))
        if cells > MAX_CELLS_WRITTEN:
            return ToolResult(ok=False, output=f"too many cells ({cells}, max {MAX_CELLS_WRITTEN})")
        try:
            target = self._target(args)
        except PathEscapeError as e:
            return ToolResult(ok=False, output=str(e))
        try:
            data = build_xlsx(sheets, args.get("chart") if isinstance(args.get("chart"), dict) else None)
        except Exception as e:  # noqa: BLE001 - bad rows or chart columns come back as a message
            return ToolResult(ok=False, output=f"couldn't build the workbook: {e}")
        target.parent.mkdir(parents=True, exist_ok=True)
        try:
            target.write_bytes(data)
        except PermissionError:
            return ToolResult(ok=False, output=f"{target.name} is open in another program (Excel?) — close it and try again")
        return ToolResult(ok=True, output=f"wrote {target.name} ({len(data)} bytes, {len(sheets)} sheet(s))")
